from pathlib import Path

from pypdf import PdfReader
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def parse_document(file_path: str) -> dict:
    """
    将不同格式文档解析为统一结构。

    返回格式：
    {
        "file_name": "...",
        "file_type": "pdf",
        "sections": [
            {
                "text": "...",
                "page": 1
            }
        ]
    }
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {file_path}")

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"暂不支持该文件类型: {suffix}"
        )

    if suffix == ".pdf":
        sections = _parse_pdf(path)

    elif suffix == ".docx":
        sections = _parse_docx(path)

    else:
        sections = _parse_text_file(path)

    # 去掉空文本
    sections = [
        section
        for section in sections
        if section["text"].strip()
    ]

    return {
        "file_name": path.name,
        "file_type": suffix.lstrip("."),
        "sections": sections
    }


def _parse_pdf(path: Path) -> list[dict]:
    """
    PDF 按页解析。

    保留 page 信息，后续引用溯源可以展示：
    xxx.pdf · 第3页
    """

    reader = PdfReader(str(path))

    sections = []

    for page_index, page in enumerate(
        reader.pages,
        start=1
    ):
        text = page.extract_text() or ""

        if text.strip():
            sections.append({
                "text": text.strip(),
                "page": page_index
            })

    return sections

def _iter_docx_blocks(document):
    """
    按 DOCX 中的原始顺序遍历顶层段落和表格。

    python-docx 的 document.paragraphs 和
    document.tables 会分别返回两类内容，
    无法保留它们原本的交错顺序。

    这里直接遍历 body XML，
    将 Paragraph / Table 按真实顺序返回。
    """

    for child in (
        document.element.body.iterchildren()
    ):
        if isinstance(child, CT_P):
            yield Paragraph(
                child,
                document
            )

        elif isinstance(child, CT_Tbl):
            yield Table(
                child,
                document
            )



def _parse_docx(path: Path) -> list[dict]:
    """
    DOCX 解析。

    按文档中的真实顺序读取：
    Paragraph -> Table -> Paragraph -> ...

    表格按行展开为纯文本，
    保持表格在原文中的相对位置。
    """

    document = Document(str(path))

    texts = []

    for block in _iter_docx_blocks(
        document
    ):
        # 普通段落
        if isinstance(
            block,
            Paragraph
        ):
            text = block.text.strip()

            if text:
                texts.append(text)

        # 表格
        elif isinstance(
            block,
            Table
        ):
            for row in block.rows:
                cells = []

                for cell in row.cells:
                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:
                        cells.append(
                            cell_text
                        )

                if cells:
                    texts.append(
                        " | ".join(cells)
                    )

    full_text = "\n".join(texts)

    if not full_text.strip():
        return []

    return [
        {
            "text": full_text,
            "page": None,
        }
    ]



def _parse_text_file(path: Path) -> list[dict]:
    """
    TXT / Markdown 文本读取。
    """

    try:
        text = path.read_text(
            encoding="utf-8"
        )

    except UnicodeDecodeError:
        text = path.read_text(
            encoding="gb18030"
        )

    if not text.strip():
        return []

    return [{
        "text": text.strip(),
        "page": None
    }]
